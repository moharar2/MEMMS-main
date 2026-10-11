"""
Unit tests for Single Equipment Maintenance Report (Passport) and Work Order Ticket export.
"""
import pytest
from src.core.database import get_db
from src.models.equipment import Equipment
from src.models.work_order import WorkOrder
from src.services.report_service import report_service


def test_equipment_history_report():
    with get_db() as db:
        eq = db.query(Equipment).first()
        assert eq is not None, "At least one equipment should exist in seed db"

        data = report_service.get_equipment_history_report_data(db, eq.id)
        assert data.get("is_single_equipment") is True
        assert data.get("equipment")["asset_number"] == eq.asset_number
        assert "work_orders" in data
        assert "pm_schedules" in data
        assert "parts_consumed" in data
        assert "summary" in data

        # Test HTML generation and RTL
        html = report_service.generate_equipment_history_html(data)
        assert "هيئة مستشفى الثورة العام النموذجي" in html
        assert eq.asset_number in html
        assert "بطاقة السجل الفني وتاريخ صيانة الجهاز الطبي" in html
        assert "توقيع" in html
        assert '<html dir="rtl" lang="ar">' in html
        assert "direction: rtl" in html
        assert 'dir="rtl"' in html


def test_work_order_ticket_report():
    with get_db() as db:
        wo = db.query(WorkOrder).first()
        assert wo is not None, "At least one work order should exist in seed db"

        data = report_service.get_work_order_ticket_data(db, wo.id)
        assert data.get("is_work_order_ticket") is True
        assert data.get("work_order")["number"] == wo.work_order_number
        assert "equipment" in data
        assert "parts" in data

        # Test HTML generation and RTL
        html = report_service.generate_work_order_ticket_html(data)
        assert "هيئة مستشفى الثورة العام النموذجي" in html
        assert wo.work_order_number in html
        assert "إشعار وتذكرة أمر صيانة طبية" in html
        assert "توقيع" in html
        assert '<html dir="rtl" lang="ar">' in html
        assert "direction: rtl" in html
        assert 'dir="rtl"' in html


def test_report_templates_share_a4_rtl_layout_and_escape_values():
    ticket = {
        "is_work_order_ticket": True,
        "work_order": {"number": "WO-RTL-01", "failure_description": "<sensor> & cable"},
        "equipment": {},
        "parts": []
    }
    ticket_html = report_service.generate_work_order_ticket_html(ticket)
    assert '<html dir="rtl" lang="ar">' in ticket_html
    assert "@page { size: A4 portrait;" in ticket_html
    assert "&lt;sensor&gt; &amp; cable" in ticket_html
    assert report_service.get_page_orientation(ticket) == "portrait"

    wide_report = {"headers": ["A", "B", "C", "D", "E", "F", "G", "H"]}
    assert report_service.get_page_orientation(wide_report) == "landscape"
    assert report_service.get_page_orientation({"headers": ["A", "B"]}) == "portrait"



def test_receiving_report_uses_shared_arabic_rtl_template():
    html = report_service.generate_receiving_report_html({
        "equipment": {"name_ar": "جهاز اختبار", "asset_number": "EQ-RTL-01"},
        "receiving": {"receipt_notes": "<ملاحظة> & بيانات"}
    })
    assert '<html dir="rtl" lang="ar">' in html
    assert "محضر استلام وتعميد جهاز طبي" in html
    assert "@page { size: A4 portrait;" in html
    assert "&lt;ملاحظة&gt; &amp; بيانات" in html
    assert "توقيع" in html



def test_reports_rtl_layout(tmp_path):
    with get_db() as db:
        eq_data = report_service.get_equipment_report_data(db)
        html = report_service.generate_html_report(eq_data)
        assert '<html dir="rtl" lang="ar">' in html
        assert "direction: rtl" in html
        assert '<table dir="rtl"' in html

        # Test Excel RTL export
        excel_path = tmp_path / "test_rtl.xlsx"
        report_service.export_to_excel(eq_data, str(excel_path))
        import openpyxl
        wb = openpyxl.load_workbook(str(excel_path))
        ws = wb.active
        assert ws.sheet_view.rightToLeft is True

        # Test prepare_rtl_document
        from PySide6.QtWidgets import QApplication
        from PySide6.QtGui import QTextDocument
        from PySide6.QtCore import Qt
        app = QApplication.instance() or QApplication([])
        doc = QTextDocument()
        report_service.prepare_rtl_document(doc, html)
        assert doc.defaultTextOption().textDirection() == Qt.LayoutDirection.RightToLeft


def test_bidi_text_protection():
    """Verify that pure LTR codes and mixed Arabic/English text are protected against BiDi inversion."""
    # Pure LTR codes and dates
    assert report_service._bidi_text("WO-THW-2026-001") == "&lrm;WO-THW-2026-001&lrm;"
    assert report_service._bidi_text("EQ-THW-013") == "&lrm;EQ-THW-013&lrm;"
    assert report_service._bidi_text("2026-09-21 19:48") == "&lrm;2026-09-21 19:48&lrm;"
    
    # Mixed Arabic + English in parentheses
    res = report_service._bidi_text("رقم الأصل (Asset Tag)")
    assert "&lrm;(Asset Tag)&lrm;" in res
    assert "رقم الأصل" in res
    
    # Mixed Arabic + English S/N
    res_sn = report_service._bidi_text("الرقم التسلسلي (S/N)")
    assert "&lrm;(S/N)&lrm;" in res_sn
    
    # Empty / None handling
    assert report_service._bidi_text(None) == "—"
    assert report_service._bidi_text("") == "—"


def test_info_table_cell_ordering_rtl():
    """Verify that info table cells are laid out so right-to-left readers see Label1 -> Val1 -> Label2 -> Val2."""
    pairs = [
        ("رقم أمر الصيانة", "WO-001"),
        ("تاريخ البلاغ", "2026-09-21")
    ]
    html = report_service._html_info_table(pairs)
    assert '<table dir="rtl" class="info-table"' in html
    
    # In the HTML row: Value 2 (left), Label 2 (mid-left), Value 1 (mid-right), Label 1 (right)
    # So the order of appearances in HTML string from left to right must be Val2 -> Lbl2 -> Val1 -> Lbl1
    pos_val2 = html.find("2026-09-21")
    pos_lbl2 = html.find("تاريخ البلاغ")
    pos_val1 = html.find("WO-001")
    pos_lbl1 = html.find("رقم أمر الصيانة")
    
    assert pos_val2 < pos_lbl2 < pos_val1 < pos_lbl1


def test_data_table_column_ordering_rtl():
    """Verify that in data tables, column 0 appears in the rightmost position in QTextTable."""
    headers = ["رقم القطعة", "اسم القطعة", "الكمية", "الإجمالي"]
    rows = [["SP-001", "حزام إغلاق", "1", "500.00"]]
    html = report_service._html_table(headers, rows)
    assert '<table dir="rtl" class="data-table"' in html
    
    # Headers in HTML string are reversed: الإجمالي -> الكمية -> اسم القطعة -> رقم القطعة
    pos_h_last = html.find("الإجمالي")
    pos_h_first = html.find("رقم القطعة")
    assert pos_h_last < pos_h_first
    
    # Row cells in HTML string are reversed: 500.00 -> 1 -> حزام إغلاق -> SP-001
    pos_c_last = html.find("500.00")
    pos_c_first = html.find("SP-001")
    assert pos_c_last < pos_c_first


def test_work_order_ticket_pdf_generation(tmp_path):
    """Verify that a full work order ticket renders into a valid, printable A4 PDF."""
    from PySide6.QtWidgets import QApplication
    from PySide6.QtGui import QTextDocument, QPdfWriter, QPageSize, QPageLayout
    from PySide6.QtCore import QSizeF, QMarginsF
    from PySide6.QtPdf import QPdfDocument
    
    app = QApplication.instance() or QApplication([])
    
    ticket_data = {
        "title": "إشعار وتذكرة أمر صيانة رقم WO-THW-2026-999",
        "is_work_order_ticket": True,
        "work_order": {
            "id": 999,
            "number": "WO-THW-2026-999",
            "created_at": "2026-09-21 19:48",
            "completed_at": "2026-09-22 14:30",
            "priority": "مرتفع",
            "status": "مكتمل",
            "reported_by": "د. أحمد علي",
            "department": "العناية المركزة",
            "technician": "م. فؤاد يحيى",
            "failure_description": "تسريب بخار في جهاز التعقيم 500L",
            "resolution_notes": "تم استبدال SP-THW-010 بنجاح",
            "labor_cost": 150.0,
            "parts_cost": 450.0,
            "total_cost": 600.0
        },
        "equipment": {
            "asset_number": "EQ-THW-013",
            "serial_number": "SN-TT-88129",
            "name_ar": "جهاز تعقيم مركزي",
            "model": "Tuttnauer HS 5090",
            "criticality": "حيوية",
            "risk_level": "عالي"
        },
        "parts": [
            {
                "part_number": "SP-THW-010",
                "name": "حزام سيليكون",
                "quantity": 1.0,
                "unit_price": 450.0,
                "total_price": 450.0
            }
        ]
    }
    
    html = report_service.generate_work_order_ticket_html(ticket_data)
    pdf_path = str(tmp_path / "test_ticket_output.pdf")
    
    writer = QPdfWriter(pdf_path)
    writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    writer.setPageOrientation(QPageLayout.Orientation.Portrait)
    writer.setPageMargins(QMarginsF(8, 8, 8, 8), QPageLayout.Unit.Millimeter)
    writer.setResolution(96)
    
    paint_rect = writer.pageLayout().paintRectPixels(writer.resolution())
    doc = QTextDocument()
    doc.setPageSize(QSizeF(paint_rect.width(), paint_rect.height()))
    report_service.prepare_rtl_document(doc, html, font_size=9)
    doc.print_(writer)
    
    # Verify PDF was created and can be opened
    pdf_doc = QPdfDocument()
    load_status = pdf_doc.load(pdf_path)
    assert pdf_doc.pageCount() >= 1


